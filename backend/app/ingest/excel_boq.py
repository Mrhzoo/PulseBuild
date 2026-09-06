"""Excel/CSV schedule and BOQ extract. Fail soft on unknown sheets."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook

from app.ingest.types import ExtractedEvent, ParseResult, detect_language

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "activity": ("activity", "task", "name", "item", "description", "وصف", "بند"),
    "start": ("start", "start_date", "start date", "from", "بداية"),
    "finish": ("finish", "finish_date", "end", "end date", "to", "نهاية"),
    "duration": ("duration", "days", "مدة"),
    "qty": ("qty", "quantity", "qty.", "كمية"),
    "unit": ("unit", "uom", "وحدة"),
    "amount": ("amount", "value", "total", "price", "قيمة", "مبلغ"),
    "item": ("item", "code", "ref", "رقم"),
}


def _norm(cell: Any) -> str:
    if cell is None:
        return ""
    return str(cell).strip()


def _map_headers(row: list[str]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    lowered = [c.strip().lower() for c in row]
    for field, aliases in HEADER_ALIASES.items():
        for i, header in enumerate(lowered):
            if header and header in aliases:
                mapped[field] = i
                break
    return mapped


def _row_event(kind, mapped, values, document_id, sheet, row_num):
    payload = {
        "sheet": sheet,
        "row": row_num,
        "pointer": f"{document_id}#sheet:{sheet}!A{row_num}",
    }
    filled = 0
    for field, idx in mapped.items():
        if idx < len(values) and values[idx]:
            payload[field] = values[idx]
            filled += 1
    if filled == 0:
        return None
    return ExtractedEvent(
        type=kind,
        payload=payload,
        confidence=0.75 if filled >= 2 else 0.55,
        pointer=payload["pointer"],
    )


def parse_csv(data: bytes, document_id: str, filename: str = "sheet.csv") -> ParseResult:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(StringIO(text)))
    if not rows:
        return ParseResult("", "needs_better_file", "unknown", [])
    mapped = _map_headers(rows[0])
    if not mapped:
        return ParseResult(
            text[:20_000],
            "parse.partial",
            detect_language(text),
            [
                ExtractedEvent(
                    type="parse.partial",
                    payload={"reason": "unknown_headers", "headers": rows[0][:20], "pointer": f"{document_id}#sheet:csv!A1"},
                    confidence=0.35,
                    pointer=f"{document_id}#sheet:csv!A1",
                )
            ],
        )
    kind = "schedule.row" if {"start", "finish", "activity"} & set(mapped) else "boq.row"
    events = []
    for i, row in enumerate(rows[1:], start=2):
        ev = _row_event(kind, mapped, [_norm(c) for c in row], document_id, "csv", i)
        if ev:
            events.append(ev)
    return ParseResult(text[:20_000], "extracted" if events else "parse.partial", detect_language(text), events)


def parse_excel(data: bytes, document_id: str, filename: str = "book.xlsx") -> ParseResult:
    if filename.lower().endswith(".csv"):
        return parse_csv(data, document_id, filename)
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return ParseResult(
            "",
            "needs_better_file",
            "unknown",
            [
                ExtractedEvent(
                    type="parse.failed",
                    payload={"reason": "unreadable_workbook", "pointer": f"{document_id}#file"},
                    confidence=0.2,
                    pointer=f"{document_id}#file",
                )
            ],
        )
    events = []
    texts = []
    unknown_sheets = []
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [_norm(c) for c in rows[0]]
        texts.append(" ".join(header))
        mapped = _map_headers(header)
        if not mapped:
            unknown_sheets.append(sheet.title)
            events.append(
                ExtractedEvent(
                    type="parse.partial",
                    payload={"reason": "unknown_sheet", "sheet": sheet.title, "headers": header[:20], "pointer": f"{document_id}#sheet:{sheet.title}!A1"},
                    confidence=0.35,
                    pointer=f"{document_id}#sheet:{sheet.title}!A1",
                )
            )
            continue
        kind = "schedule.row" if {"start", "finish", "activity"} & set(mapped) else "boq.row"
        for i, raw in enumerate(rows[1:], start=2):
            values = [_norm(c) for c in raw]
            texts.append(" ".join(values))
            ev = _row_event(kind, mapped, values, document_id, sheet.title, i)
            if ev:
                events.append(ev)
    joined = "\n".join(texts)
    status = "extracted"
    if unknown_sheets:
        status = "parse.partial"
    if not events:
        status = "needs_better_file"
    return ParseResult(joined[:20_000], status, detect_language(joined), events)
